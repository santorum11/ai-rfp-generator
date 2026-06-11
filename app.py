import streamlit as st
from rfp_crew import generate_clarifying_questions, draft_rfp_and_get_roles, calculate_budget_and_review
from fpdf import FPDF
import io

# Import text parsing and creation libraries
import pypdf
import docx2txt
import openpyxl
from pptx import Presentation
from pptx.util import Inches, Pt
import docx

st.set_page_config(page_title="Agentic RFP Builder", layout="wide")

# --- Helper Function: Universal File Text Extractor ---
def extract_text_from_file(uploaded_file):
    file_name = uploaded_file.name.lower()
    extracted_text = ""
    if file_name.endswith(('.txt', '.md')):
        extracted_text = uploaded_file.read().decode("utf-8")
    elif file_name.endswith('.pdf'):
        reader = pypdf.PdfReader(uploaded_file)
        for page in reader.pages:
            extracted_text += (page.extract_text() or "") + "\n"
    elif file_name.endswith('.docx'):
        extracted_text = docx2txt.process(uploaded_file)
    elif file_name.endswith('.xlsx'):
        wb = openpyxl.load_workbook(uploaded_file, data_only=True)
        for sheet in wb.sheetnames:
            ws = wb[sheet]
            extracted_text += f"\n--- Sheet: {sheet} ---\n"
            for row in ws.iter_rows(values_only=True):
                row_text = " ".join([str(cell) for cell in row if cell is not None])
                if row_text.strip():
                    extracted_text += row_text + "\n"
    elif file_name.endswith('.pptx'):
        prs = Presentation(uploaded_file)
        for i, slide in enumerate(prs.slides):
            extracted_text += f"\n--- Slide {i+1} ---\n"
            for shape in slide.shapes:
                if hasattr(shape, "text") and shape.text.strip():
                    extracted_text += shape.text + "\n"
    return extracted_text.strip()

# --- Helper Function: Export PDF ---
def create_pdf(text):
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("helvetica", size=11)
    clean_text = text.encode('latin-1', 'replace').decode('latin-1')
    pdf.multi_cell(0, 6, clean_text)
    return bytes(pdf.output())

# --- Helper Function: Export Word Document (.docx) ---
def create_docx(text):
    doc = docx.Document()
    clean_text = "".join(ch for ch in text if ord(ch) >= 32 or ch in "\n\r\t")
    for line in clean_text.split('\n'):
        if line.startswith('# '):
            doc.add_heading(line.replace('# ', ''), level=1)
        elif line.startswith('## '):
            doc.add_heading(line.replace('## ', ''), level=2)
        elif line.startswith('### '):
            doc.add_heading(line.replace('### ', ''), level=3)
        else:
            doc.add_paragraph(line)
    bio = io.BytesIO()
    doc.save(bio)
    return bio.getvalue()

# --- Helper Function: Export PowerPoint Presentation (.pptx) ---
def create_pptx(text):
    prs = Presentation()
    title_slide_layout = prs.slide_layouts[0]
    slide = prs.slides.add_slide(title_slide_layout)
    slide.shapes.title.text = "Project RFP & Proposal"
    slide.placeholders[1].text = "Generated autonomously by Agentic RFP Builder"

    bullet_slide_layout = prs.slide_layouts[1]
    lines = text.split('\n')
    current_tf = None

    for line in lines:
        line = line.strip()
        if not line: continue
        if line.startswith('# ') or line.startswith('## ') or line.startswith('### '):
            slide = prs.slides.add_slide(bullet_slide_layout)
            slide.shapes.title.text = line.lstrip('#').strip()
            current_tf = slide.placeholders[1].text_frame
            current_tf.word_wrap = True
        elif current_tf:
            p = current_tf.add_paragraph()
            p.text = line.lstrip('- ').lstrip('* ').strip()
            p.font.size = Pt(14)
            
    bio = io.BytesIO()
    prs.save(bio)
    return bio.getvalue()

# --- Helper Function: Markdown Table to CSV ---
def markdown_to_csv(markdown_text):
    csv_lines = []
    for line in markdown_text.split('\n'):
        if '|' in line:
            clean_line = line.strip().strip('|')
            csv_lines.append(",".join([col.strip() for col in clean_line.split('|')]))
    return "\n".join(csv_lines)

# --- 🧠 SESSION STATE MEMORY ---
if "step" not in st.session_state:
    st.session_state.step = 1

st.title("🚀 Agentic RFP & Budget Estimation Platform")

# ==========================================
# STEP 1: INITIAL INPUT
# ==========================================
if st.session_state.step == 1:
    col1, col2 = st.columns(2)
    with col1:
        company_name = st.text_input("Client / Company Name:", placeholder="e.g., Acme Logistics")
    
    st.markdown("---")
    st.markdown("### 📝 Project Scope Input")
    st.info("Provide your project requirements by using either Option A OR Option B below.")

    input_col, file_col = st.columns(2)
    
    with input_col:
        st.markdown("#### Option A: Type Requirements")
        raw_input = st.text_area("Enter raw project scope or rough notes:", height=180, placeholder="Type here...")
    with file_col:
        st.markdown("#### Option B: Upload Requirements File")
        uploaded_file = st.file_uploader(
            "Upload any brief format (.txt, .md, .pdf, .docx, .xlsx, .pptx)", 
            type=["txt", "md", "pdf", "docx", "xlsx", "pptx"]
        )
        if uploaded_file is not None:
            st.success(f"📎 Attached file successfully: {uploaded_file.name}")

    st.markdown("---")

    if st.button("Step 1: Start Discovery Phase", type="primary"):
        final_requirements = ""
        if raw_input.strip():
            final_requirements = raw_input.strip()
        elif uploaded_file is not None:
            with st.spinner("Extracting content from document..."):
                final_requirements = extract_text_from_file(uploaded_file)
        
        if company_name.strip() and final_requirements:
            with st.status("🚀 Initializing Agent Framework...", expanded=True) as status_box:
                questions = generate_clarifying_questions(
                    company_name, 
                    final_requirements, 
                    status_callback=status_box.write
                )
                status_box.update(label="Discovery Analysis Complete!", state="complete", expanded=False)
                
            st.session_state.company_name = company_name
            st.session_state.raw_input = final_requirements
            st.session_state.questions = questions
            st.session_state.step = 2
            st.rerun()
        else:
            st.error("Missing Information: Please ensure you provide a Company Name and either type requirements OR upload a supported file.")

# ==========================================
# STEP 2: DISCOVERY (QA)
# ==========================================
elif st.session_state.step == 2:
    st.info("The AI needs a few more details to generate an accurate RFP.")
    st.markdown(f"### 🔍 Clarifying Questions\n{st.session_state.questions}")
    
    user_answers = st.text_area("Provide your answers here:", height=150)
    
    col1, col2 = st.columns([1, 4])
    with col1:
        if st.button("Step 2: Draft RFP", type="primary"):
            if user_answers.strip():
                with st.status("🚀 Launching Multi-Agent Drafting Sequence...", expanded=True) as status_box:
                    brief, rfp, roles = draft_rfp_and_get_roles(
                        st.session_state.company_name, 
                        st.session_state.raw_input, 
                        user_answers,
                        status_callback=status_box.write
                    )
                    status_box.update(label="RFP Document Architecture Formed!", state="complete", expanded=False)
                    
                st.session_state.client_brief = brief
                st.session_state.rfp_text = rfp
                # Using a generic "Rate" key here allows the dynamic header configuration in Step 3
                st.session_state.role_data = [{"Role": r, "Rate": 50} for r in roles]
                st.session_state.step = 3
                st.rerun()
            else:
                st.error("Please provide answers to continue.")
    with col2:
        if st.button("Start Over"):
            st.session_state.step = 1
            st.rerun()

# ==========================================
# STEP 3: DYNAMIC ROLE & CURRENCY EDITOR
# ==========================================
elif st.session_state.step == 3:
    st.success("✨ RFP Drafted!")
    st.markdown("### 💸 Adjust Team, Currency & Rates")
    
    # Currency configuration relocated to Step 3 layout flow with CAD added
    col_curr, _ = st.columns([3, 5])
    with col_curr:
        currency_selection = st.selectbox(
            "Select Project Estimation Currency:",
            ["USD ($)", "INR (₹)", "CAD ($)", "EUR (€)", "GBP (£)"]
        )
        # Isolate target raw character symbol
        st.session_state.currency_symbol = currency_selection.split("(")[1].replace(")", "")
    
    st.info(f"Edit rates below in ({st.session_state.currency_symbol}), append target responsibilities, or clear rows.")
    
    # Interactive Data Editor mapping
    edited_roles = st.data_editor(
        st.session_state.role_data,
        num_rows="dynamic",
        column_config={
            "Role": st.column_config.TextColumn("Role Required", required=True),
            "Rate": st.column_config.NumberColumn(f"Hourly Rate ({st.session_state.currency_symbol})", min_value=1, required=True)
        },
        use_container_width=True
    )

    col1, col2 = st.columns([1, 4])
    with col1:
        if st.button("Step 3: Calculate Final Budget", type="primary"):
            custom_rates = {row["Role"]: row["Rate"] for row in edited_roles if row.get("Role")}
            
            with st.status("🚀 Launching Finance & Auditing Crew...", expanded=True) as status_box:
                final_out, budget_text = calculate_budget_and_review(
                    st.session_state.company_name,
                    st.session_state.client_brief,
                    st.session_state.rfp_text,
                    custom_rates,
                    st.session_state.currency_symbol,
                    status_callback=status_box.write
                )
                status_box.update(label="Financial Modeling Complete!", state="complete", expanded=False)
                
            st.session_state.budget_text = budget_text
            st.session_state.final_out = final_out
            st.session_state.step = 4
            st.rerun()
    with col2:
        if st.button("Start Over"):
            st.session_state.step = 1
            st.rerun()

# ==========================================
# STEP 4: FINAL OUTPUT
# ==========================================
elif st.session_state.step == 4:
    st.success("✅ Multi-Agent Workflow Complete!")
    st.markdown("### 💾 Export Documents")
    
    col_format, col_rfp_btn, col_budget_btn, col_reset = st.columns([2, 2, 2, 2])
    
    with col_format:
        selected_format = st.selectbox(
            "Select RFP Format:", 
            ["PDF (.pdf)", "Word Document (.docx)", "PowerPoint Presentation (.pptx)", "Markdown (.md)", "Plain Text (.txt)"]
        )
    
    if selected_format == "PDF (.pdf)":
        rfp_data = create_pdf(st.session_state.rfp_text)
        file_ext = "pdf"
        mime_type = "application/pdf"
    elif selected_format == "Word Document (.docx)":
        rfp_data = create_docx(st.session_state.rfp_text)
        file_ext = "docx"
        mime_type = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    elif selected_format == "PowerPoint Presentation (.pptx)":
        rfp_data = create_pptx(st.session_state.rfp_text)
        file_ext = "pptx"
        mime_type = "application/vnd.openxmlformats-officedocument.presentationml.presentation"
    elif selected_format == "Markdown (.md)":
        rfp_data = st.session_state.rfp_text.encode("utf-8")
        file_ext = "md"
        mime_type = "text/markdown"
    else:
        rfp_data = st.session_state.rfp_text.encode("utf-8")
        file_ext = "txt"
        mime_type = "text/plain"

    with col_rfp_btn:
        st.write(" ")
        st.write(" ")
        st.download_button(
            label=f"📄 Download RFP ({file_ext.upper()})",
            data=rfp_data,
            file_name=f"{st.session_state.company_name}_RFP.{file_ext}",
            mime=mime_type,
            type="primary"
        )
        
    with col_budget_btn:
        st.write(" ")
        st.write(" ")
        csv_data = markdown_to_csv(st.session_state.budget_text)
        st.download_button(
            label="📊 Download Budget (CSV)", 
            data=csv_data, 
            file_name=f"{st.session_state.company_name}_Budget.csv", 
            mime="text/csv"
        )
        
    with col_reset:
        st.write(" ")
        st.write(" ")
        if st.button("Start New Project"):
            st.session_state.step = 1
            st.rerun()
            
    st.markdown("---")
    st.markdown(st.session_state.final_out)